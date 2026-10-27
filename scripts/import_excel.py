"""Read the workbook without modifying it; regenerate the versioned source dataset."""
import sys,json,collections,datetime
import openpyxl
w=openpyxl.load_workbook(sys.argv[1],data_only=True)
def clean(v):
 return str(v).strip() if v is not None else ''
def num(v):
 return v if isinstance(v,(int,float)) else 0
records=[]
s=w['Data']
for row in s.iter_rows(min_row=3,values_only=True):
 if not isinstance(row[0],(int,float)) or not row[1]:continue
 records.append(dict(sourceRow=int(row[0])+2,sourceNotes=['']*12,id='all-'+str(int(row[0])),scope='all',tag=clean(row[2]),description=clean(row[1]),workshop=clean(row[3]),type=clean(row[4]),frequency=clean(row[5]),remarks=clean(row[6]),owner=clean(row[7]),plan=[num(row[8+m]) for m in range(12)],actual=[num(row[21+m]) for m in range(12)],errors=[None]*12))
s=w['SSCC Master Plan']
for i in range(3,s.max_row+1):
 r=[s.cell(i,c).value for c in range(1,29)]
 if not isinstance(r[0],(int,float)) or not r[2]:continue
 a=[s.cell(i+1,11+m).value for m in range(12)]
 records.append(dict(sourceRow=i,sourceNotes=[v if isinstance(v,str) else '' for v in a],id='critical-'+str(i),scope='critical',tag=clean(r[2]),description=clean(r[1]),workshop=clean(r[3]),type=clean(r[4]),frequency=clean(r[5]),remarks=clean(r[6]),owner=clean(r[7]),plan=[num(r[10+m]) for m in range(12)],actual=[int(isinstance(v,(int,float))) for v in a],errors=[v if isinstance(v,(int,float)) else None for v in a]))
result=dict(year=2026,source='Calibration critical instrument plant 2_2026.xlsx',reportMonth=8,records=records,notes=['Data: 387 instruments (2 missing TAGs); monthly source plan totals 726; cached KPI / pivot totals 724 (November 183 vs 181). Dashboard recalculates from source rows.','Critical: 111 instruments from SSCC Master Plan. Plan 327, actual 154. Actual counts numeric cells (Excel COUNT), including zero and negative readings; excludes 8 No cal. cells. Cached summary 304/145 omits rows after 188; readings are retained separately and are not summed as job counts.','Data and critical sheets have different workshop/type/frequency assignments. Scopes remain separate; no silent merge by TAG.','Workbook dashboard selects September. Later months with no actual entries mean no recorded result, not a confirmed failure. No tolerance limits or production tonnage are supplied.'])
with open('data/source.json','w') as f:json.dump(result,f,ensure_ascii=False,indent=2)
for scope in ['all','critical']:
 rs=[r for r in records if r['scope']==scope]
 print(scope,len(rs),'plan',sum(sum(r['plan']) for r in rs),'actual',sum(sum(r['actual']) for r in rs))
